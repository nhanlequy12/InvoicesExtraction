import csv
import io
import json
from decimal import Decimal
from typing import List
from src.core.logging import logger
from src.normalization.schemas import NormalizedInvoice


class InvoiceFileExporter:
    """
    Exports normalized invoices to various standard financial formats:
    - Excel (.xlsx) Purchase Ledger / Bang ke hoa don mua vao
    - JSON structured data
    - CSV
    """

    def export_to_json(self, invoices: List[NormalizedInvoice]) -> str:
        data = [inv.model_dump(mode="json") for inv in invoices]
        return json.dumps(data, indent=2, ensure_ascii=False)

    def export_to_csv(self, invoices: List[NormalizedInvoice]) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        # Write headers
        writer.writerow([
            "Invoice Series",
            "Invoice Number",
            "Issue Date",
            "Seller Tax ID",
            "Seller Name",
            "Buyer Name",
            "Subtotal",
            "VAT Amount",
            "Total Amount",
            "Currency",
            "Signature Status"
        ])

        for inv in invoices:
            writer.writerow([
                inv.metadata.invoice_series,
                inv.metadata.invoice_number,
                inv.metadata.issue_date,
                inv.seller.tax_id or "",
                inv.seller.name,
                inv.buyer.name,
                str(inv.summary.subtotal),
                str(inv.summary.vat_amount),
                str(inv.summary.total_amount),
                inv.summary.currency,
                inv.signature.status
            ])

        return output.getvalue()

    def export_to_excel(self, invoices: List[NormalizedInvoice]) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bảng kê Hóa đơn"

        # Headers
        headers = [
            "STT", "Ký hiệu HĐ", "Số HĐ", "Ngày lập", "MST Người bán",
            "Tên người bán", "Tiền chưa thuế", "Tiền thuế GTGT", "Tổng tiền thanh toán",
            "Loại tiền", "Chữ ký số", "Trạng thái"
        ]
        ws.append(headers)

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

        for col_idx, cell in enumerate(ws[1], start=1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, inv in enumerate(invoices, start=1):
            ws.append([
                idx,
                inv.metadata.invoice_series,
                inv.metadata.invoice_number,
                inv.metadata.issue_date,
                inv.seller.tax_id or "",
                inv.seller.name,
                float(inv.summary.subtotal),
                float(inv.summary.vat_amount),
                float(inv.summary.total_amount),
                inv.summary.currency,
                inv.signature.status,
                "Đã duyệt"
            ])

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
